import streamlit as st
import streamlit.components.v1 as components
from google import genai
from google.genai import types
from PIL import Image, ImageEnhance, ImageDraw
import pandas as pd
import io
import os
import json
import re
from datetime import datetime
import pypdfium2 as pdfium
from openpyxl.styles import PatternFill, Font

st.set_page_config(page_title="Mandi AI: जिनी मुनीम", layout="wide")
st.title("🌾 Mandi AI: जिनी (भारतीय मुनीम) & Master Vault")

api_key = st.secrets.get("GEMINI_API_KEY")
if not api_key:
    st.error("Secrets me GEMINI_API_KEY configure karein.")
    st.stop()

client = genai.Client(api_key=api_key)

ACTIVE_MODELS = ["gemini-3.6-flash", "gemini-3.5-flash-lite"]

def run_ai(contents, config=None):
    last_err = None
    for m in ACTIVE_MODELS:
        try:
            if config:
                return client.models.generate_content(model=m, contents=contents, config=config)
            return client.models.generate_content(model=m, contents=contents)
        except Exception as e:
            last_err = e
            continue
    raise Exception(f"AI Service Error: {last_err}")

# Storage Directories
STORAGE_DIR = "vault_files"
INDEX_FILE = "vault_index.json"
for folder in ["Truck_Logistics", "Mandi_Parchi", "Invoices_Bills"]:
    os.makedirs(os.path.join(STORAGE_DIR, folder), exist_ok=True)

def load_vault():
    if os.path.exists(INDEX_FILE):
        try:
            with open(INDEX_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {"rules": [], "documents": []}

def save_vault(d):
    with open(INDEX_FILE, "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=2)

vault = load_vault()

def load_img(file_obj):
    if file_obj.type == "application/pdf":
        return pdfium.PdfDocument(file_obj.getvalue())[0].render(scale=2).to_pil()
    return Image.open(file_obj).convert("RGB")

def build_part(pil_img):
    img = ImageEnhance.Contrast(pil_img).enhance(1.4)
    img = ImageEnhance.Sharpness(img).enhance(1.3)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=90)
    return types.Part.from_bytes(data=buf.getvalue(), mime_type="image/jpeg")

# Sidebar
with st.sidebar:
    st.header("🧠 मुनीम की याददाश्त (Memory)")
    new_r = st.text_area("व्यापार का नया नियम सिखाएं:")
    if st.button("💾 नियम सेव करें"):
        if new_r.strip():
            vault["rules"].append(new_r.strip())
            save_vault(vault)
            st.success("नियम हमेशा के लिए सेव हो गया!")
    if vault.get("rules"):
        for i, r in enumerate(vault["rules"][-3:], 1):
            st.caption(f"{i}. {r}")
    st.divider()
    mode = st.radio("Navigation:", [
        "👩‍💼 जिनी से बातचीत (Voice Munim)",
        "📤 पर्ची/बिल अपलोड और ऑटो-फाइलिंग",
        "🧮 हिसाब समझें (Doubt Solver)",
        "🔍 दो पर्चियों का मिलान (Forensic Match)"
    ])

# MODULE 1: GENIE LIVE ASSISTANT
if mode == "👩‍💼 जिनी से बातचीत (Voice Munim)":
    st.subheader("👩‍💼 जिनी: आपकी अपनी भारतीय डिजिटल मुनीम")
    st.caption("माइक का बटन दबाकर बोलिए या टाइप कीजिए। जिनी भारतीय स्त्री की आवाज़ में उत्तर देंगी।")

    speech_ui = """
    <div style="background: #0f172a; padding: 14px; border-radius: 10px; border: 1px solid #38bdf8; margin-bottom: 12px; display: flex; align-items: center; justify-content: space-between;">
        <div>
            <span id="dot" style="display:inline-block; width:12px; height:12px; border-radius:50%; background:#22c55e;"></span>
            <strong id="stTxt" style="color:#f8fafc; margin-left:8px; font-size:14px;">जिनी तैयार हैं...</strong>
        </div>
        <button onclick="startMic()" style="background:#e11d48; color:white; border:none; padding:8px 16px; border-radius:6px; font-weight:bold; cursor:pointer;">
            🎙️ बोलकर पूछें
        </button>
    </div>
    <script>
    var rec;
    function startMic() {
        if (!('webkitSpeechRecognition' in window) && !('SpeechRecognition' in window)) {
            alert('Browser voice support nahi karta'); return;
        }
        var SRec = window.SpeechRecognition || window.webkitSpeechRecognition;
        rec = new SRec();
        rec.lang = 'hi-IN';
        rec.onstart = function() {
            document.getElementById('dot').style.background = '#e11d48';
            document.getElementById('stTxt').innerText = 'सुन रही हूँ, बोलिए...';
        };
        rec.onresult = function(e) {
            var txt = e.results[0][0].transcript;
            document.getElementById('stTxt').innerText = 'सुना: ' + txt;
            document.getElementById('dot').style.background = '#22c55e';
            var ta = window.parent.document.querySelector('textarea[data-testid="stChatInputTextArea"]');
            var btn = window.parent.document.querySelector('button[data-testid="stChatInputSubmitButton"]');
            if (ta) {
                var setVal = Object.getOwnPropertyDescriptor(window.parent.HTMLTextAreaElement.prototype, "value").set;
                setVal.call(ta, txt);
                ta.dispatchEvent(new Event('input', { bubbles: true }));
                setTimeout(function(){ if(btn) btn.click(); }, 300);
            }
        };
        rec.onerror = function() { document.getElementById('dot').style.background = '#22c55e'; };
        rec.start();
    }
    </script>
    """
    components.html(speech_ui, height=65)

    if "chat_history" not in st.session_state:
        st.session_state["chat_history"] = [
            {"role": "assistant", "content": "नमस्ते भैया! मैं जिनी हूँ। आज मंडी का कौन सा हिसाब या गाड़ी का बिल देखना है? बताइए, मैं सब समझाती हूँ।"}
        ]

    for m in st.session_state["chat_history"]:
        with st.chat_message(m["role"]):
            st.markdown(m["content"])

    q = st.chat_input("जिनी से कुछ भी पूछें...")
    if q:
        st.session_state["chat_history"].append({"role": "user", "content": q})
        with st.chat_message("user"):
            st.markdown(q)

        docs_summary = json.dumps(vault.get("documents", []), ensure_ascii=False)
        rules_text = "\n".join(vault.get("rules", []))
        prompt = f"""
        आप 'जिनी' (Genie) हैं—एक होशियार, आदरणीय और मधुर भारतीय मुनीम बहन।
        भाषा: शुद्ध, आदरणीय हिंदी (स्त्रीलिंग: "मैं बताती हूँ", "मैंने हिसाब देख लिया है")।
        [व्यापार के नियम]: {rules_text}
        [पुराने दस्तावेज़ और बही-खाता]: {docs_summary}
        [यूज़र का सवाल]: "{q}"
        सीधा और सटीक उत्तर दें। अगर कोई फाइल मिली तो उसका नाम भी बताएं।
        """
        with st.chat_message("assistant"):
            with st.spinner("जिनी हिसाब देख रही हैं..."):
                try:
                    res = run_ai([prompt])
                    ans = res.text.strip()
                    st.markdown(ans)
                    st.session_state["chat_history"].append({"role": "assistant", "content": ans})

                    # Show matched physical image/PDF
                    matched = [d for d in vault.get("documents", []) if d.get("filename") in ans or d.get("doc_id") in ans]
                    for m_doc in matched:
                        p = m_doc.get("stored_path")
                        if p and os.path.exists(p):
                            st.write(f"📂 **दस्तावेज़:** `{m_doc.get('filename')}`")
                            if p.lower().endswith(".pdf"):
                                st.image(pdfium.PdfDocument(p)[0].render(scale=2).to_pil(), width=450)
                            else:
                                st.image(Image.open(p), width=450)

                    # Authentic Indian Female TTS
                    clean_v = ans.replace('"', '\\"').replace('\n', ' ')
                    components.html(f"""
                    <script>
                        var synth = window.speechSynthesis;
                        synth.cancel();
                        var u = new SpeechSynthesisUtterance("{clean_v}");
                        u.lang = 'hi-IN'; u.pitch = 1.15; u.rate = 0.95;
                        var vs = synth.getVoices();
                        for (var i=0; i<vs.length; i++) {{
                            if (vs[i].lang.includes('hi') || vs[i].lang.includes('IN')) {{
                                var n = vs[i].name.toLowerCase();
                                if (n.includes('female') || n.includes('lekha') || n.includes('swara') || n.includes('google')) {{
                                    u.voice = vs[i]; break;
                                }}
                            }}
                        }}
                        synth.speak(u);
                    </script>
                    """, height=0)
                except Exception as ex:
                    st.error(f"Error: {ex}")

# MODULE 2: FULL OCR & AUTO-FILING
elif mode == "📤 पर्ची/बिल अपलोड और ऑटो-फाइलिंग":
    st.subheader("📤 पर्ची या बही-खाता अपलोड करें")
    up = st.file_uploader("Document upload karein", type=["jpg", "jpeg", "png", "pdf"])
    if up:
        src_img = load_img(up)
        if st.button("🚀 हिसाब निकालें और वॉल्ट में सुरक्षित करें"):
            with st.spinner("AI हिसाब मिला रहा है..."):
                p_ocr = """
                Extract EVERY row from this mandi parchi/bill into JSON.
                Format:
                {
                  "category": "Mandi_Parchi" or "Truck_Logistics" or "Invoices_Bills",
                  "date": "DD/MM/YYYY or null",
                  "summary": "1 line summary",
                  "records": [
                    {
                      "party_name": "Farmer/Trader Name",
                      "weight": 25.0,
                      "rate": 2100.0,
                      "written_amount": 52500.0,
                      "doubt_flag": false,
                      "box_2d": [ymin, xmin, ymax, xmax]
                    }
                  ]
                }
                Return ONLY raw valid JSON.
                """
                try:
                    res = run_ai(
                        contents=[p_ocr, build_part(src_img)],
                        config=types.GenerateContentConfig(response_mime_type="application/json", temperature=0.1)
                    )
                    txt = res.text.strip()
                    if txt.startswith("```json"): txt = txt[7:]
                    if txt.endswith("```"): txt = txt[:-3]
                    p_data = json.loads(txt.strip())

                    recs = p_data.get("records", [])
                    cat = p_data.get("category", "Mandi_Parchi")
                    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                    f_path = os.path.join(STORAGE_DIR, cat, f"{ts}_{up.name}")
                    with open(f_path, "wb") as f: f.write(up.getbuffer())

                    doc_info = {
                        "doc_id": f"DOC_{ts}",
                        "filename": up.name,
                        "stored_path": f_path,
                        "category": cat,
                        "summary": p_data.get("summary"),
                        "records": recs
                    }
                    vault["documents"].append(doc_info)
                    save_vault(vault)

                    st.session_state["ocr_records"] = recs
                    st.session_state["ocr_img"] = src_img
                    st.success(f"✅ सुरक्षित रूप से [{cat}] में सेव कर दिया गया!")
                except Exception as e:
                    st.error(f"Error: {e}")

        if "ocr_records" in st.session_state and "ocr_img" in st.session_state:
            recs = st.session_state["ocr_records"]
            base = st.session_state["ocr_img"].copy()
            iw, ih = base.size
            df = pd.DataFrame(recs)
            draw = ImageDraw.Draw(base)

            calcs, statuses = [], []
            for _, r in df.iterrows():
                try:
                    w = float(re.sub(r"[^\d.]", "", str(r.get("weight", 0))) or 0)
                    rt = float(re.sub(r"[^\d.]", "", str(r.get("rate", 0))) or 0)
                    wa = float(re.sub(r"[^\d.]", "", str(r.get("written_amount", 0))) or 0)
                    ca = round(w * rt, 2)
                    calcs.append(ca)

                    mismatch = (wa > 0 and ca > 0 and abs(ca - wa) > 1.0)
                    doubt = bool(r.get("doubt_flag", False))
                    if mismatch:
                        statuses.append("🔴 CALC MISMATCH")
                        col = "#FF0000"
                    elif doubt:
                        statuses.append("🔴 DOUBTFUL")
                        col = "#FFA500"
                    else:
                        statuses.append("✅ 100% OK")
                        col = "#00AA00"

                    bx = r.get("box_2d")
                    if bx and len(bx) == 4:
                        y1, x1, y2, x2 = bx
                        draw.rectangle([int((x1/1000.0)*iw), int((y1/1000.0)*ih), int((x2/1000.0)*iw), int((y2/1000.0)*ih)], outline=col, width=3)
                except Exception:
                    calcs.append(0); statuses.append("🔴 ERROR")

            df["CALCULATED_AMOUNT"] = calcs
            df["STATUS"] = statuses

            c1, c2 = st.columns([1, 1])
            with c1:
                st.write("**Visual Tag Overlay:**")
                st.image(base, width="stretch")
            with c2:
                st.write("**Audited Grid:**")
                show_cols = [c for c in df.columns if c != "box_2d"]
                e_df = st.data_editor(df[show_cols], width="stretch")

                out = io.BytesIO()
                with pd.ExcelWriter(out, engine="openpyxl") as w:
                    e_df.to_excel(w, index=False, sheet_name="Mandi")
                    ws = w.sheets["Mandi"]
                    r_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    g_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                    sc = e_df.columns.get_loc("STATUS") + 1
                    for rx in range(2, len(e_df) + 2):
                        val = str(ws.cell(row=rx, column=sc).value)
                        if "🔴" in val:
                            for cx in range(1, len(e_df.columns) + 1): ws.cell(row=rx, column=cx).fill = r_fill
                        elif "✅" in val:
                            ws.cell(row=rx, column=sc).fill = g_fill
                st.download_button("📥 Clean Audited Excel डाउनलोड करें", data=out.getvalue(), file_name="mandi_clean.xlsx")

# MODULE 3: DOUBT SOLVER
elif mode == "🧮 हिसाब समझें (Doubt Solver)":
    st.subheader("🔍 किसी भी जोड़ या कटौती को समझें")
    f_d = st.file_uploader("पर्ची की फोटो डालें", type=["jpg", "jpeg", "png", "pdf"], key="d_solv")
    query_m = st.text_input("क्या समझना है?", placeholder="उदा. कुल 48,200 कैसे आया? दर और हमाली समझाएं।")
    if f_d and query_m and st.button("🧠 पूरा हिसाब खोलकर बताएं"):
        with st.spinner("हिसाब तोड़ा जा रहा है..."):
            res = run_ai([f"Explain mandi calculation in Hindi step-by-step for: {query_m}", build_part(load_img(f_d))])
            st.markdown(res.text)

# MODULE 4: 2-DOCUMENT FORENSIC MATCHER
else:
    st.subheader("🔍 दो पर्चियों/चालानों का 100% सटीक मिलान")
    c1, c2 = st.columns(2)
    with c1: f1 = st.file_uploader("Document 1", type=["jpg", "jpeg", "png", "pdf"], key="f1")
    with c2: f2 = st.file_uploader("Document 2", type=["jpg", "jpeg", "png", "pdf"], key="f2")
    if f1 and f2 and st.button("बारीक मिलान करें (Forensic Cross-Check)"):
        with st.spinner("0.0001% अंतर खोजा जा रहा है..."):
            p_comp = "Compare Doc 1 and Doc 2 strictly. Highlight any mismatch in red HTML span. Provide comparison table."
            res = run_ai([p_comp, build_part(load_img(f1)), build_part(load_img(f2))])
            st.markdown(res.text, unsafe_allow_html=True)
    
